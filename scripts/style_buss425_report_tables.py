from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "outputs" / "buss425" / "report" / "buss425_final_report.docx"

HEADER_FILL = "0F4C5C"
ALT_FILL = "F4F8FA"
WHITE = "FFFFFF"
BORDER = "B8C7CE"
TEXT = "18212B"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, *, bold: bool = False, color: str = TEXT, size: float = 8.8) -> None:
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(size)
            run.font.bold = bold
            color_element = run._element.get_or_add_rPr().find(qn("w:color"))
            if color_element is None:
                color_element = OxmlElement("w:color")
                run._element.get_or_add_rPr().append(color_element)
            color_element.set(qn("w:val"), color)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def add_front_matter_page_breaks(document) -> None:
    paragraphs = document.paragraphs
    if paragraphs and paragraphs[0].style.name == "Title":
        remove_paragraph(paragraphs[0])

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Course: BUSS425"):
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        if paragraph.text.startswith("Abstract\n1. Introduction"):
            paragraph.add_run().add_break(WD_BREAK.PAGE)


def keep_headings_with_following_text(document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text == "Responsible Use Statement":
            paragraph.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)


def main() -> None:
    document = Document(DOCX_PATH)
    add_front_matter_page_breaks(document)
    keep_headings_with_following_text(document)
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if row_index == 0:
                    set_cell_shading(cell, HEADER_FILL)
                    set_cell_text(cell, bold=True, color=WHITE)
                else:
                    set_cell_shading(cell, ALT_FILL if row_index % 2 == 0 else WHITE)
                    set_cell_text(cell)

    document.save(DOCX_PATH)
    print(f"Styled and centered {len(document.tables)} tables in {DOCX_PATH}.")


if __name__ == "__main__":
    main()
